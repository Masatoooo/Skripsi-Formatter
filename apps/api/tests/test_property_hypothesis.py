from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Pt, Twips
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st

from app.core.fixers.runner import fix_runner
from app.core.integrity import normalized_text
from app.core.rules.evaluator import make_finding_id
from app.core.rules.loader import profile_loader
from app.core.units import (
    cm_to_emu,
    cm_to_twips,
    emu_to_cm,
    half_points_to_pt,
    line_spacing_multiplier_to_wline,
    pt_to_half_points,
    pt_to_twips,
    twips_to_cm,
    twips_to_pt,
    wline_to_line_spacing_multiplier,
)


@given(cm=st.floats(min_value=0.5, max_value=30.0, allow_nan=False, allow_infinity=False))
def test_property_cm_twips_and_emu_roundtrip(cm: float):
    tw = cm_to_twips(cm)
    back_cm = twips_to_cm(tw)
    assert abs(back_cm - cm) <= 0.01

    emu = cm_to_emu(cm)
    back_emu_cm = emu_to_cm(emu)
    assert abs(back_emu_cm - cm) <= 1e-4


@given(pt=st.integers(min_value=8, max_value=36).map(float))
def test_property_pt_roundtrip(pt: float):
    assert twips_to_pt(pt_to_twips(pt)) == pt
    assert half_points_to_pt(pt_to_half_points(pt)) == pt


@given(mult=st.sampled_from([1.0, 1.15, 1.5, 2.0]))
def test_property_line_spacing_roundtrip(mult: float):
    wline = line_spacing_multiplier_to_wline(mult)
    assert abs(wline_to_line_spacing_multiplier(wline) - mult) <= 1e-3


@given(
    rule_id=st.sampled_from(["PAGE-002", "FONT-001", "PARA-001", "HEAD-001"]),
    sec=st.integers(min_value=0, max_value=10),
    blk=st.integers(min_value=0, max_value=500),
    actual=st.text(min_size=1, max_size=30),
)
def test_property_finding_id_stability(rule_id: str, sec: int, blk: int, actual: str):
    id1 = make_finding_id(rule_id, sec, blk, actual)
    id2 = make_finding_id(rule_id, sec, blk, actual)
    assert id1 == id2
    assert id1.startswith(f"f-{rule_id.lower()}-")


@settings(
    max_examples=12,
    deadline=None,
    suppress_health_check=[HealthCheck.function_scoped_fixture],
)
@given(
    top_cm=st.floats(min_value=1.5, max_value=5.0),
    left_cm=st.floats(min_value=1.5, max_value=5.0),
    font_name=st.sampled_from(["Arial", "Calibri", "Times New Roman", "Courier New"]),
    font_pt=st.sampled_from([10, 11, 12, 14]),
    para_texts=st.lists(
        st.text(
            alphabet=st.characters(whitelist_categories=("L", "N"), whitelist_characters=" .,;:"),
            min_size=10,
            max_size=80,
        ).filter(lambda s: len(s.strip()) >= 5),
        min_size=1,
        max_size=4,
    ),
)
def test_property_synthetic_docx_fix_integrity_and_idempotence(
    top_cm: float,
    left_cm: float,
    font_name: str,
    font_pt: int,
    para_texts: list[str],
):
    profile = profile_loader.load("contoh-s1-skripsi")
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        in_doc_path = tmp / "in.docx"
        out_1_path = tmp / "out1.docx"
        out_2_path = tmp / "out2.docx"

        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Twips(cm_to_twips(top_cm))
        sect.left_margin = Twips(cm_to_twips(left_cm))

        for t in para_texts:
            p = doc.add_paragraph()
            r = p.add_run(t)
            r.font.name = font_name
            r.font.size = Pt(font_pt)

        doc.save(str(in_doc_path))

        # 1st fix run
        _, _ = fix_runner.run(in_doc_path, out_1_path, profile, mode="all")
        assert normalized_text(in_doc_path) == normalized_text(out_1_path)

        # 2nd fix run (idempotence)
        changes_2, _ = fix_runner.run(out_1_path, out_2_path, profile, mode="all")
        assert changes_2 == []
        assert normalized_text(out_1_path) == normalized_text(out_2_path)
