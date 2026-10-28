from __future__ import annotations

import io
import logging
import zipfile
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.config import settings
from app.core.messages import ErrorCode
from app.main import app, reset_rate_limiter
from app.services.storage import storage
from tests.fixture_builder import ensure_all_fixtures

client = TestClient(app)


@pytest.fixture(autouse=True)
def _clean_rate_limit():
    reset_rate_limiter()
    yield
    reset_rate_limiter()


def test_reject_invalid_extensions_and_macro_docm():
    fixtures_dir = ensure_all_fixtures()
    macro_bytes = (fixtures_dir / "macro.docm").read_bytes()

    res_docm = client.post(
        "/v1/documents",
        files={"file": ("skripsi.docm", macro_bytes, "application/octet-stream")},
    )
    assert res_docm.status_code == 400
    assert res_docm.headers.get("X-Error-Code") == ErrorCode.INGEST_MACRO_DETECTED

    res_doc = client.post(
        "/v1/documents",
        files={"file": ("skripsi.doc", b"dummy legacy doc", "application/msword")},
    )
    assert res_doc.status_code == 400
    assert res_doc.headers.get("X-Error-Code") == ErrorCode.INGEST_LEGACY_DOC

    res_pdf = client.post(
        "/v1/documents",
        files={"file": ("skripsi.pdf", b"%PDF-1.4", "application/pdf")},
    )
    assert res_pdf.status_code == 400
    assert res_pdf.headers.get("X-Error-Code") == ErrorCode.INGEST_INVALID_EXT


def test_reject_empty_and_corrupt_zip():
    res_empty = client.post(
        "/v1/documents",
        files={"file": ("kosong.docx", b"", "application/octet-stream")},
    )
    assert res_empty.status_code == 400
    assert res_empty.headers.get("X-Error-Code") == ErrorCode.INGEST_EMPTY_FILE

    res_bad_zip = client.post(
        "/v1/documents",
        files={"file": ("rusak.docx", b"not a zip archive", "application/octet-stream")},
    )
    assert res_bad_zip.status_code == 400
    assert res_bad_zip.headers.get("X-Error-Code") == ErrorCode.INGEST_INVALID_ZIP


def test_reject_xxe_traversal_and_zip_bomb():
    fixtures_dir = ensure_all_fixtures()

    # 1. XXE
    xxe_bytes = (fixtures_dir / "bad_xxe.docx").read_bytes()
    res_xxe = client.post(
        "/v1/documents",
        files={"file": ("bad_xxe.docx", xxe_bytes, "application/octet-stream")},
    )
    assert res_xxe.status_code == 400
    assert res_xxe.headers.get("X-Error-Code") == ErrorCode.INGEST_XXE_DETECTED

    # 2. Path traversal
    trav_bytes = (fixtures_dir / "bad_traversal.docx").read_bytes()
    res_trav = client.post(
        "/v1/documents",
        files={"file": ("bad_traversal.docx", trav_bytes, "application/octet-stream")},
    )
    assert res_trav.status_code == 400
    assert res_trav.headers.get("X-Error-Code") == ErrorCode.INGEST_PATH_TRAVERSAL

    # 3. Zip bomb (too many entries)
    bomb_bytes = (fixtures_dir / "bad_zip_bomb.docx").read_bytes()
    res_bomb = client.post(
        "/v1/documents",
        files={"file": ("bad_zip_bomb.docx", bomb_bytes, "application/octet-stream")},
    )
    assert res_bomb.status_code == 400
    assert res_bomb.headers.get("X-Error-Code") == ErrorCode.INGEST_TOO_MANY_ENTRIES


def test_reject_embedded_vbaproject_bin():
    doc = Document()
    doc.add_paragraph("Test macro inside .docx")
    buf = io.BytesIO()
    doc.save(buf)

    out_buf = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(buf.getvalue()), "r") as zin,
        zipfile.ZipFile(out_buf, "w") as zout,
    ):
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("word/vbaProject.bin", b"fake macro binary")

    res = client.post(
        "/v1/documents",
        files={"file": ("stealth_macro.docx", out_buf.getvalue(), "application/octet-stream")},
    )
    assert res.status_code == 400
    assert res.headers.get("X-Error-Code") == ErrorCode.INGEST_MACRO_DETECTED


def test_storage_ttl_cleanup_and_explicit_delete(tmp_path: Path):
    fixtures_dir = ensure_all_fixtures()
    clean_bytes = (fixtures_dir / "clean_styled.docx").read_bytes()

    up_res = client.post(
        "/v1/documents",
        files={"file": ("clean.docx", clean_bytes, "application/octet-stream")},
    )
    assert up_res.status_code == 201
    doc_id = up_res.json()["doc_id"]
    assert storage.session_exists(doc_id)

    del_res = client.delete(f"/v1/documents/{doc_id}")
    assert del_res.status_code == 200
    assert not storage.session_exists(doc_id)

    # Test TTL expiration cleanup
    doc_id_2 = storage.create_document_session("temp.docx")
    storage.save_upload(doc_id_2, clean_bytes)
    assert storage.session_exists(doc_id_2)
    deleted_count = storage.cleanup_expired(max_age_seconds=-1)
    assert deleted_count >= 1
    assert not storage.session_exists(doc_id_2)


def test_no_document_content_in_logs(caplog: pytest.LogCaptureFixture):
    secret_marker = "RAHASIA_SKRIPSI_MAHASISWA_99887766"
    doc = Document()
    doc.add_paragraph(
        f"Kalimat ini memuat penanda unik {secret_marker} yang tidak boleh bocor ke log."
    )
    buf = io.BytesIO()
    doc.save(buf)

    with caplog.at_level(logging.INFO):
        up_res = client.post(
            "/v1/documents",
            files={"file": ("privasi.docx", buf.getvalue(), "application/octet-stream")},
        )
        assert up_res.status_code == 201
        doc_id = up_res.json()["doc_id"]

        an_res = client.post(
            f"/v1/documents/{doc_id}/analyze",
            json={"profile_id": "contoh-s1-skripsi"},
        )
        assert an_res.status_code == 200

        fx_res = client.post(
            f"/v1/documents/{doc_id}/fix",
            json={"profile_id": "contoh-s1-skripsi", "mode": "all"},
        )
        assert fx_res.status_code == 200
        client.delete(f"/v1/documents/{doc_id}")

    for record in caplog.records:
        assert secret_marker not in record.getMessage()


def test_rate_limiter_enforced(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(settings, "rate_limit_per_minute", 3)
    reset_rate_limiter()

    for _ in range(3):
        r = client.get("/health")
        assert r.status_code == 200

    r_blocked = client.get("/health")
    assert r_blocked.status_code == 429
    assert r_blocked.headers.get("X-Error-Code") == ErrorCode.RATE_LIMIT_EXCEEDED
